#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_excel.py
=============
Turn qr_details.csv (or any CSV from these tools) into a proper .xlsx workbook.

    python make_excel.py                      # pdf_files/qr_details.csv -> .xlsx
    python make_excel.py --csv other.csv --out report.xlsx
    python make_excel.py --csv pdf_files/manifest.csv

Why not just rename the CSV: Excel mangles this data on open. Saudi mobile
numbers (0512345678) lose their leading zero and 10-digit company numbers can
flip to scientific notation. This writes those columns as real text cells so
they survive, sets the sheet right-to-left for Arabic, and adds a frozen
filtered header.
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is missing.  Install it with:\n"
             "    .venv/bin/pip install openpyxl")

HERE = Path(__file__).resolve().parent
DEFAULT_CSV = HERE / "pdf_files" / "qr_details.csv"

# Columns that must stay text, or Excel eats them:
#   phone/fax  -> leading zero stripped  (0512345678 becomes 512345678)
#   uid/cert_no-> 10 digits may render as 7.04201E+09
TEXT_COLUMNS = {"phone", "fax", "uid", "cert_no", "mobile"}

# Nicer Arabic/English headers for the known columns; anything unknown is
# passed through unchanged, so this also works on manifest.csv.
HEADERS = {
    "pdf_file":            "الملف / File",
    "cert_no":             "رقم الشهادة / Cert no.",
    "uid":                 "الرقم الوطني / Company no.",
    "facility_name":       "اسم المنشأة / Facility",
    "owner_name":          "اسم المالك / Owner",
    "phone":               "رقم الهاتف / Phone",
    "email":               "البريد الإلكتروني / Email",
    "city":                "المدينة / City",
    "region":              "المنطقة / Region",
    "activity":            "النشاط / Activity",
    "fax":                 "الفاكس / Fax",
    "case_id":             "رقم الحالة / Case ID",
    "classification_type": "نوع التصنيف / Type",
    "qr_url":              "رابط QR / QR link",
    # manifest.csv
    "page":   "الصفحة / Page",
    "name":   "اسم المنشأة / Facility",
    "expiry": "تاريخ الانتهاء / Expiry",
    "status": "الحالة / Status",
    "type":   "النوع / Type",
    "file":   "الملف / File",
}

WIDTHS = {  # sensible starting widths; everything else is auto-sized
    "facility_name": 46, "owner_name": 30, "activity": 46,
    "email": 30, "qr_url": 52, "pdf_file": 44, "name": 46,
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", default=str(DEFAULT_CSV), help="input CSV")
    ap.add_argument("--out", default=None, help="output .xlsx (default: same name)")
    ap.add_argument("--sheet", default="البيانات", help="worksheet name")
    ap.add_argument("--ltr", action="store_true", help="left-to-right sheet")
    args = ap.parse_args()

    src = Path(args.csv).expanduser().resolve()
    if not src.is_file():
        sys.exit(f"No such CSV: {src}")
    dst = Path(args.out).expanduser().resolve() if args.out else src.with_suffix(".xlsx")

    with src.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        sys.exit(f"{src} has no data rows")
    cols = list(rows[0].keys())

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet[:31]
    ws.sheet_view.rightToLeft = not args.ltr      # Arabic reads right-to-left

    # --- header ---
    head_fill = PatternFill("solid", fgColor="1F6F5C")
    head_font = Font(bold=True, color="FFFFFF", size=11)
    for c, key in enumerate(cols, 1):
        cell = ws.cell(row=1, column=c, value=HEADERS.get(key, key))
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    # --- data ---
    link_font = Font(color="0563C1", underline="single")
    for r, row in enumerate(rows, 2):
        for c, key in enumerate(cols, 1):
            val = (row.get(key) or "").strip()
            cell = ws.cell(row=r, column=c)
            if key in TEXT_COLUMNS:
                cell.value = val
                cell.number_format = "@"          # force text, keep leading zeros
                cell.alignment = Alignment(horizontal="left")
            elif key == "qr_url" and val.startswith("http"):
                cell.value = val
                cell.hyperlink = val
                cell.font = link_font
            else:
                cell.value = val
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

    # --- widths ---
    for c, key in enumerate(cols, 1):
        if key in WIDTHS:
            width = WIDTHS[key]
        else:
            longest = max([len(str(HEADERS.get(key, key)))]
                          + [len((r.get(key) or "")) for r in rows[:400]])
            width = min(max(longest + 2, 10), 50)
        ws.column_dimensions[get_column_letter(c)].width = width

    wb.save(dst)

    size_kb = dst.stat().st_size / 1024
    print(f"Wrote {dst}")
    print(f"  {len(rows)} rows x {len(cols)} columns, {size_kb:.0f} KB")
    filled = sum(1 for r in rows if (r.get("phone") or "").strip())
    if "phone" in cols:
        print(f"  {filled}/{len(rows)} rows have a phone number")
    return 0


if __name__ == "__main__":
    sys.exit(main())
